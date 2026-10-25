from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase, SimpleTestCase, override_settings
from openpyxl import load_workbook

from finance_app.transactions.models import CardMovement, CardStatement, CardPaymentAllocation, CreditCard, MovementType, Transaction
from finance_app.transactions.services.card_statements import parse_card_pages, merchant_category, store_card_statement
from finance_app.transactions.services.card_reconciliation import allocate_payment, auto_reconcile, bank_remaining, card_remaining


class CardParserTests(SimpleTestCase):
    def test_old_card_uses_original_purchase_even_when_already_paid(self):
        text='''Tarjeta de Crédito # 1111 2222 3333 1234
Banco Davivienda
Periodo liquidado ENE.01 /26-ENE.31 /26
Saldo anterior $0
Consumos mes $100
Pagos Mes $100
Saldo Total $0
0001234 20260102 APPLE.COM/BILL $100 $0+ $0 $0+ 1 / 1 20.00 2.00
0003456 20260103 PAGO $100 $100- $0 $0+ 0 / 0 0.00 0.00'''
        parsed=parse_card_pages([text])
        self.assertEqual(parsed['rows'][0]['amount'],Decimal('100'))
        self.assertEqual(parsed['rows'][0]['billed_amount'],Decimal('0'))
        self.assertEqual(parsed['rows'][1]['amount'],Decimal('-100'))
        self.assertEqual(parsed['validation_notes'],[])

    def test_merchant_categories_and_viva_is_not_iva(self):
        from finance_app.transactions.services.card_statements import movement_kind
        self.assertEqual(merchant_category('APPLE.COM/BILL'),('Apple','tecnologia'))
        self.assertEqual(merchant_category('AMZN Mktp US'),('Amazon','compras'))
        self.assertEqual(merchant_category('MERCADOLIBRE*TIENDA'),('Mercado Libre','compras'))
        self.assertEqual(merchant_category('DLO*DIDI FOOD CO PAYIN'),('DiDi Food','restaurantes'))
        self.assertEqual(merchant_category('UNKNOWN MERCHANT')[1],'otros')
        self.assertEqual(movement_kind('VIVA 1 A IPS'),'purchase')
        self.assertEqual(movement_kind('IVA SERVICIOS'),'tax')

    def test_reversed_card_payment_is_a_charge_not_an_abono(self):
        text='''Tarjeta de Crédito # 1111 2222 3333 1234
Banco Davivienda
Periodo liquidado ENE.01 /26-ENE.31 /26
Saldo anterior $0
Consumos mes $0
Pagos Mes $0
Saldo Total $100
0001234 20260102 PAGO $100 $100+ $0 $0+ 0 / 0 0.00 0.00'''
        parsed=parse_card_pages([text])
        self.assertEqual(parsed['rows'][0]['kind'],'adjustment')
        self.assertEqual(parsed['rows'][0]['amount'],100)


class CardLedgerTests(TestCase):
    def setUp(self):
        self.user=get_user_model().objects.create_user(username='card-user',is_superuser=True,is_staff=True)
        self.other=get_user_model().objects.create_user(username='other-card-user')
        self.client.force_login(self.user)
        self.card=CreditCard.objects.create(user=self.user,bank='Davivienda',last4='1234')
        self.kind=MovementType.objects.get_or_create(name='pago tarjeta credito')[0]

    def payment(self,amount=100,when=date(2026,2,2),fingerprint='payment'):
        return CardMovement.objects.create(card=self.card,date=when,kind='payment',amount=-Decimal(amount),description='PAGO',fingerprint=fingerprint)

    def bank(self,amount=100,when=date(2026,1,31),description='Pago Tarj. Credito N1111222233331234'):
        return Transaction.objects.create(user=self.user,date=when,amount=-Decimal(amount),currency='COP',movement_type=self.kind,description=description,classification='gasto')

    def test_unique_cross_month_payment_matches_and_is_idempotent(self):
        tx=self.bank();payment=self.payment()
        self.assertEqual(auto_reconcile(self.user),1)
        self.assertEqual(auto_reconcile(self.user),0)
        self.assertEqual(bank_remaining(tx),0);self.assertEqual(card_remaining(payment),0)
        self.assertEqual(Transaction.objects.count(),1)

    def test_ambiguous_payments_and_wrong_card_are_not_matched(self):
        self.bank();self.payment();self.payment(fingerprint='second')
        self.assertEqual(auto_reconcile(self.user),0)
        CardMovement.objects.all().delete()
        self.payment();Transaction.objects.update(description='Pago Tarj. Credito N1111222233339999')
        self.assertEqual(auto_reconcile(self.user),0)

    def test_partial_allocations_enforce_budgets_and_owner(self):
        tx=self.bank(150);first=self.payment();second=self.payment(fingerprint='second')
        allocate_payment(self.user,tx.pk,first.pk,100)
        with self.assertRaises(ValueError):allocate_payment(self.user,tx.pk,second.pk,60)
        allocate_payment(self.user,tx.pk,second.pk,50)
        self.assertEqual(card_remaining(second),50)
        with self.assertRaises(Transaction.DoesNotExist):allocate_payment(self.other,tx.pk,second.pk,1)
        self.assertEqual(CardPaymentAllocation.objects.count(),2)

    def test_multiple_expenses_can_share_a_payment_without_exceeding_budget(self):
        from finance_app.transactions.services.card_reconciliation import allocate_expenses,expense_remaining,bank_expense_remaining
        from finance_app.transactions.models import CardExpenseAllocation
        tx=self.bank(150)
        first=CardMovement.objects.create(card=self.card,date=date(2026,1,1),kind='purchase',amount=100,description='Apple',fingerprint='expense1')
        second=CardMovement.objects.create(card=self.card,date=date(2026,1,2),kind='purchase',amount=100,description='Amazon',fingerprint='expense2')
        amount,count=allocate_expenses(self.user,tx.pk,[first.pk,second.pk])
        self.assertEqual((amount,count),(150,2))
        self.assertEqual(expense_remaining(first),0)
        self.assertEqual(expense_remaining(second),50)
        self.assertEqual(bank_expense_remaining(tx),0)
        with self.assertRaises(ValueError):allocate_expenses(self.user,tx.pk,[second.pk],1)
        self.assertEqual(CardExpenseAllocation.objects.count(),2)
        # The expense breakdown is independent of the statement payment match.
        payment=self.payment(150)
        allocate_payment(self.user,tx.pk,payment.pk,150)
        self.assertEqual(card_remaining(payment),0)

    def test_expense_selection_rejects_other_users_and_post_accepts_multiple_ids(self):
        from finance_app.transactions.models import CardExpenseAllocation
        tx=self.bank(100)
        first=CardMovement.objects.create(card=self.card,date=date(2026,1,1),kind='purchase',amount=60,description='Apple',fingerprint='expense1')
        second=CardMovement.objects.create(card=self.card,date=date(2026,1,2),kind='purchase',amount=40,description='Amazon',fingerprint='expense2')
        response=self.client.post('/cards/reconcile/',{'action':'expenses','bank_id':tx.pk,'expense_ids':[first.pk,second.pk]})
        self.assertEqual(response.status_code,302)
        self.assertEqual(CardExpenseAllocation.objects.count(),2)
        self.client.force_login(self.other)
        self.client.post('/cards/reconcile/',{'action':'expenses','bank_id':tx.pk,'expense_ids':[first.pk]})
        self.assertEqual(CardExpenseAllocation.objects.count(),2)

    def test_backup_exports_contain_valid_database_and_referenced_pdfs(self):
        import sqlite3,zipfile,json
        with TemporaryDirectory() as root,override_settings(MEDIA_ROOT=root):
            source=Path(root)/'source.sqlite3'
            with sqlite3.connect(source) as connection:
                connection.execute('create table marker (value text)')
                connection.execute("insert into marker values ('snapshot')")
            tx=self.bank();tx.source_pdf='original.pdf';tx.save()
            folder=Path(root)/'pdfs';folder.mkdir();(folder/'original.pdf').write_bytes(b'%PDF-test')
            with patch('finance_app.core.views._backup_source_path',return_value=source):
                response=self.client.get('/backup/download/')
                database=b''.join(response.streaming_content);response.close()
                copied=Path(root)/'export.sqlite3';copied.write_bytes(database)
                with sqlite3.connect(copied) as connection:
                    self.assertEqual(connection.execute('pragma integrity_check').fetchone()[0],'ok')
                    self.assertEqual(connection.execute('select value from marker').fetchone()[0],'snapshot')
                response=self.client.get('/backup/full/download/')
                archive=BytesIO(b''.join(response.streaming_content));response.close()
            with zipfile.ZipFile(archive) as bundle:
                self.assertIsNone(bundle.testzip())
                self.assertEqual(bundle.read('uploads/pdfs/original.pdf'),b'%PDF-test')
                self.assertIn('db.sqlite3',bundle.namelist())
                self.assertEqual(json.loads(bundle.read('manifest.json'))['missing_pdfs'],[])
        self.client.force_login(self.other)
        self.assertEqual(self.client.get('/backup/full/download/').status_code,302)

    def test_statement_duplicates_and_repeated_installments_do_not_duplicate_spend(self):
        parsed=dict(bank='Davivienda',last4='1234',period_start=date(2026,1,1),period_end=date(2026,1,31),currency='COP',validation_notes=[],rows=[dict(date=date(2026,1,2),description='APPLE.COM/BILL',amount=Decimal('100'),kind='purchase',reference='ref1',billed_amount=Decimal('25'),source_page=1)])
        with TemporaryDirectory() as root,override_settings(MEDIA_ROOT=root):
            path=Path(root)/'sample.pdf';path.write_bytes(b'example')
            first,created=store_card_statement(self.user,path,parsed)
            self.assertTrue(created)
            _,created=store_card_statement(self.user,path,parsed);self.assertFalse(created)
            parsed.update(period_start=date(2026,2,1),period_end=date(2026,2,28))
            store_card_statement(self.user,path,parsed)
        self.assertEqual(CardMovement.objects.count(),1)
        self.assertEqual(CardStatement.objects.count(),2)
        response=self.client.get('/cards/')
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.context['purchase_total'],100)
        self.assertEqual(Transaction.objects.count(),0)

    def test_card_workbook_is_filtered_private_and_has_real_cells(self):
        CardMovement.objects.create(card=self.card,date=date(2026,1,2),kind='purchase',amount=100,description='=1+1',merchant='Apple',category='tecnologia',fingerprint='p1')
        CardMovement.objects.create(card=self.card,date=date(2026,2,2),kind='purchase',amount=200,description='other month',fingerprint='p2')
        response=self.client.get('/cards/?month=2026-01&export=1')
        self.assertEqual(response.status_code,200)
        book=load_workbook(BytesIO(response.content))
        self.assertEqual(book['Consumos'].max_row,2)
        self.assertEqual(book['Consumos']['E2'].value,'=1+1')
        self.assertEqual(book['Consumos']['E2'].data_type,'s')
        self.assertEqual(book['Consumos']['I2'].value,100)
        self.assertEqual(len(book.sheetnames),8)
        self.client.force_login(self.other)
        book=load_workbook(BytesIO(self.client.get('/cards/?export=1').content))
        self.assertEqual(book['Consumos'].max_row,1)

    def test_bank_export_empty_and_filtered_and_reference_leading_zeros(self):
        response=self.client.get('/transactions/?export=1')
        book=load_workbook(BytesIO(response.content));self.assertEqual(book.active.max_row,1)
        tx=self.bank();tx.reference='00001234';tx.description='=1+1';tx.save()
        response=self.client.get('/transactions/?export=1&classification=gasto')
        book=load_workbook(BytesIO(response.content));headers=[c.value for c in book.active[1]]
        values={headers[i]:book.active.cell(2,i+1) for i in range(len(headers))}
        self.assertEqual(values['reference'].value,'00001234')
        self.assertEqual(values['description'].data_type,'s')
        self.assertEqual(values['amount'].value,-100)
        self.assertEqual(values['transaction_id'].value,tx.pk)
        book=load_workbook(BytesIO(self.client.get('/transactions/?export=1&classification=ingreso').content))
        self.assertEqual(book.active.max_row,1)

    def test_card_actions_are_owner_scoped_and_require_post(self):
        payment=self.payment()
        self.assertEqual(self.client.get('/cards/reconcile/').status_code,405)
        self.client.force_login(self.other)
        self.assertEqual(self.client.post(f'/cards/movements/{payment.pk}/classify/',{'category':'compras','merchant':'test'}).status_code,404)

    def test_classifying_an_expense_returns_to_the_filtered_cards_page(self):
        expense=CardMovement.objects.create(card=self.card,date=date(2026,1,2),kind='purchase',amount=100,description='APPLE.COM/BILL',fingerprint='expense')
        response=self.client.post(f'/cards/movements/{expense.pk}/classify/',{'category':'tecnologia','merchant':'Apple','next':'/cards/?month=2026-01'})
        self.assertRedirects(response,'/cards/?month=2026-01')
        expense.refresh_from_db()
        self.assertEqual(expense.category,'tecnologia')
        self.assertEqual(expense.merchant,'Apple')
        self.assertTrue(expense.manually_classified)
